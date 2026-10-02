import docx
import re
import os
from docx.shared import Inches

def md_to_docx(md_path, docx_path):
    doc = docx.Document()
    
    with open(md_path, 'r', encoding='utf-8') as f:
        text = f.read()
    
    lines = text.split('\n')
    
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
            
        # Images: ![Alt Text](path/to/image.png)
        img_match = re.match(r'^!\[(.*?)\]\((.*?)\)$', line)
        if img_match:
            alt_text = img_match.group(1)
            img_path = img_match.group(2)
            
            if os.path.exists(img_path):
                try:
                    doc.add_picture(img_path, width=Inches(5))
                    p = doc.add_paragraph(alt_text)
                    p.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
                except Exception as e:
                    doc.add_paragraph(f"[Failed to load image: {img_path}]")
            else:
                doc.add_paragraph(f"[Image not found: {img_path}]")
            i += 1
            continue
            
        # Tables (Very basic markdown table parsing)
        if line.startswith('|') and line.endswith('|'):
            # Count columns
            cols = [c.strip() for c in line.split('|') if c]
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                # Skip divider line (e.g., |---|---|)
                if '---' not in lines[i]:
                    row = [c.strip() for c in lines[i].strip().split('|') if c]
                    # if len(row) == len(cols): # removing this check to be more forgiving
                    table_lines.append(row)
                i += 1
            
            if table_lines:
                table = doc.add_table(rows=len(table_lines), cols=len(cols))
                table.style = 'Table Grid'
                for r_idx, row_data in enumerate(table_lines):
                    for c_idx, cell_data in enumerate(row_data):
                        if c_idx < len(cols):
                            table.cell(r_idx, c_idx).text = cell_data
            continue
            
        # Headings
        if line.startswith('# '):
            doc.add_heading(line[2:].strip(), level=1)
        elif line.startswith('## '):
            doc.add_heading(line[3:].strip(), level=2)
        elif line.startswith('### '):
            doc.add_heading(line[4:].strip(), level=3)
        elif line.startswith('#### '):
            doc.add_heading(line[5:].strip(), level=4)
        elif line.startswith('- '):
            p = doc.add_paragraph(style='List Bullet')
            parse_inline(p, line[2:])
        elif re.match(r'^\d+\.\s', line):
            p = doc.add_paragraph(style='List Number')
            content = re.sub(r'^\d+\.\s', '', line)
            parse_inline(p, content)
        else:
            p = doc.add_paragraph()
            parse_inline(p, line)
            
        i += 1

    doc.save(docx_path)
    print(f"Successfully converted {md_path} to {docx_path}")

def parse_inline(p, text):
    # Regex to capture **bold** or *italic*
    tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*)', text)
    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**'):
            run = p.add_run(token[2:-2])
            run.bold = True
        elif token.startswith('*') and token.endswith('*'):
            run = p.add_run(token[1:-1])
            run.italic = True
        else:
            p.add_run(token)

if __name__ == "__main__":
    md_file = "manuscript.md"
    docx_file = "Epigenetic_OS_Cancer_Plasticity_Manuscript_Revised.docx"
    md_to_docx(md_file, docx_file)
